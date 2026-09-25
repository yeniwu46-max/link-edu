"""Turn uploaded files into ordered structural blocks (heading / text / table) with page numbers."""
import io
import re
from dataclasses import dataclass, field
from pathlib import PurePath

from rag.cleaning import clean_line, is_noise, strip_page_furniture


class ParseError(ValueError):
    pass


@dataclass
class Block:
    text: str
    kind: str = 'text'  # heading | text | table
    level: int = 0
    page: int | None = None
    page_end: int | None = None
    rows: list = field(default_factory=list)


@dataclass
class ParsedDocument:
    blocks: list
    page_count: int | None = None


_NUM = '一二三四五六七八九十百零〇两'
HEADING_PATTERNS = (
    (re.compile(rf'^第[{_NUM}\d]+(?:编|篇|部分)(?!课)'), 1),
    (re.compile(rf'^第[{_NUM}\d]+章'), 2),
    (re.compile(rf'^第[{_NUM}\d]+节(?!课)'), 3),
    (re.compile(rf'^[{_NUM}]+[、．.]\s*\S'), 4),
    (re.compile(rf'^[（(][{_NUM}]+[)）]\s*\S'), 5),
    (re.compile(r'^\d+(?:\.\d+)+\s*[^\d.\s]'), None),  # 1.1 / 1.2.3 -> depth-based
    (re.compile(r'^\d+[、．.]\s*\S'), 7),
    (re.compile(r'^[（(]\d+[)）]\s*\S'), 8),
)
_SENTENCE_END = '。！？!?；;…'
_CLOSERS = '”’」』）)》'
_LIST_START = re.compile(rf'^(?:[-*•·●○■□▪]\s|\d+[、．.)]|[（(]\d+[)）]|[{_NUM}]+[、．.]|[（(][{_NUM}]+[)）])')
_MD_HEADING = re.compile(r'^(#{1,6})\s+(.+?)\s*#*\s*$')
_MD_IMAGE = re.compile(r'!\[([^\]]*)\]\([^)]*\)')
_MD_LINK = re.compile(r'\[([^\]]+)\]\([^)]*\)')
_MD_EMPHASIS = re.compile(r'(\*\*|\*|`)(?=\S)(.+?)(?<=\S)\1')


def detect_file_type(filename):
    suffix = PurePath(filename or '').suffix.lower().lstrip('.')
    mapping = {'pdf': 'pdf', 'docx': 'docx', 'txt': 'txt', 'md': 'md', 'markdown': 'md'}
    if suffix == 'doc':
        raise ParseError('暂不支持旧版 .doc 文件，请另存为 .docx 后上传')
    if suffix not in mapping:
        raise ParseError('仅支持 PDF、Word(.docx)、TXT、Markdown 文件')
    return mapping[suffix]


def heading_level(line):
    """Level for numbered headings such as 第一章 / 一、 / （一） / 1.2; None for body text."""
    stripped = line.strip()
    if not stripped or len(stripped) > 40 or stripped[-1] in '。；;！!？?，,':
        return None
    for pattern, level in HEADING_PATTERNS:
        if pattern.match(stripped):
            if level is None:
                return min(9, 5 + stripped.split()[0].count('.'))
            return level
    return None


def _ends_sentence(line):
    tail = line.rstrip(_CLOSERS) or line
    return tail[-1:] in _SENTENCE_END or tail[-1:] in '：:'


def _join(left, right):
    if left and right and left[-1].isascii() and left[-1].isalnum() and right[0].isascii() and right[0].isalnum():
        return f'{left} {right}'
    return left + right


def decode_text(data):
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode('utf-8', errors='replace')


def parse_txt(data):
    blocks = []
    for raw in decode_text(data).replace('\r\n', '\n').replace('\r', '\n').split('\n'):
        line = clean_line(raw)
        if not line or is_noise(line):
            continue
        level = heading_level(line)
        blocks.append(Block(line, 'heading', level) if level else Block(line))
    return ParsedDocument(blocks)


def _md_inline(text):
    text = _MD_IMAGE.sub(r'\1', text)
    text = _MD_LINK.sub(r'\1', text)
    return _MD_EMPHASIS.sub(r'\2', text)


def parse_markdown(data):
    blocks, paragraph, table, code = [], [], [], []
    in_code = False

    def flush_paragraph():
        if paragraph:
            text = '\n'.join(clean_line(_md_inline(line)) for line in _join_lines(paragraph))
            level = heading_level(text) if '\n' not in text else None
            if level:
                blocks.append(Block(text, 'heading', level))
            elif text and not is_noise(text):
                blocks.append(Block(text))
            paragraph.clear()

    def flush_table():
        rows = [row for row in table if not re.fullmatch(r'[|\s:\-]+', row)]
        if rows:
            cells = [' | '.join(c.strip() for c in row.strip().strip('|').split('|')) for row in rows]
            blocks.append(Block('\n'.join(cells), 'table', rows=cells))
        table.clear()

    for raw in decode_text(data).replace('\r\n', '\n').split('\n'):
        if raw.strip().startswith('```'):
            flush_paragraph()
            if in_code and any(line.strip() for line in code):
                blocks.append(Block('\n'.join(code).strip('\n')))
            code.clear()
            in_code = not in_code
            continue
        if in_code:
            code.append(raw.rstrip())
            continue
        line = raw.strip()
        if line.startswith('|'):
            flush_paragraph()
            table.append(_md_inline(line))
            continue
        flush_table()
        match = _MD_HEADING.match(line)
        if match:
            flush_paragraph()
            blocks.append(Block(clean_line(_md_inline(match.group(2))), 'heading', len(match.group(1))))
        elif not line or re.fullmatch(r'[-*_]{3,}', line):
            flush_paragraph()
        else:
            line = line.lstrip('>').strip()
            if paragraph and _LIST_START.match(line):
                flush_paragraph()
            paragraph.append(line)
    flush_paragraph()
    flush_table()
    return ParsedDocument(blocks)


def _join_lines(lines):
    """Rejoin visually wrapped lines; keep list items on their own line."""
    merged = []
    for line in lines:
        if merged and not _LIST_START.match(line):
            merged[-1] = _join(merged[-1], line)
        else:
            merged.append(line)
    return merged


def pdf_pages_to_blocks(pages):
    """pages: list of raw page texts. Paragraphs are rebuilt from wrapped lines and may span pages."""
    cleaned = strip_page_furniture([str(text or '').split('\n') for text in pages])
    blocks, current = [], None
    for number, lines in enumerate(cleaned, start=1):
        for line in lines:
            level = heading_level(line)
            if level:
                if current:
                    blocks.append(current)
                    current = None
                blocks.append(Block(line, 'heading', level, number, number))
                continue
            starts_new = current is None or _ends_sentence(current.text) or bool(_LIST_START.match(line))
            if starts_new:
                if current:
                    blocks.append(current)
                current = Block(line, page=number, page_end=number)
            else:
                current.text = _join(current.text, line)
                current.page_end = number
    if current:
        blocks.append(current)
    return blocks


def parse_pdf(data):
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt('')
        pages = [page.extract_text() or '' for page in reader.pages]
    except ParseError:
        raise
    except Exception as error:
        raise ParseError('PDF 文件无法解析，请确认文件未损坏且未加密') from error
    blocks = pdf_pages_to_blocks(pages)
    if not any(block.kind != 'heading' for block in blocks):
        raise ParseError('未能从 PDF 中提取文字，可能是扫描件，请先做 OCR 后上传')
    return ParsedDocument(blocks, len(pages))


def _docx_heading_level(paragraph):
    style = (paragraph.style.name if paragraph.style is not None else '') or ''
    match = re.match(r'^(?:Heading|标题)\s*(\d)$', style.strip(), re.I)
    if match:
        return int(match.group(1))
    if style.strip().lower() in ('title', '标题'):
        return 1
    ppr = paragraph._p.pPr
    outline = ppr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl') \
        if ppr is not None else None
    if outline is not None:
        value = outline.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
        if value is not None and value.isdigit() and int(value) < 9:
            return int(value) + 1
    return None


def _docx_table_rows(table):
    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            text = clean_line(cell.text.replace('\n', ' '))
            # Merged cells are repeated by python-docx for every grid column they span.
            if text and (not cells or cells[-1] != text):
                cells.append(text)
        if cells:
            rows.append(' | '.join(cells))
    return rows


def parse_docx(data):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(io.BytesIO(data))
    except Exception as error:
        raise ParseError('Word 文件无法解析，请确认是有效的 .docx 文件') from error
    blocks = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            paragraph = Paragraph(child, document)
            text = clean_line(paragraph.text)
            if not text or is_noise(text):
                continue
            level = _docx_heading_level(paragraph) or heading_level(text)
            blocks.append(Block(text, 'heading', level) if level else Block(text))
        elif tag == 'tbl':
            rows = _docx_table_rows(Table(child, document))
            if rows:
                blocks.append(Block('\n'.join(rows), 'table', rows=rows))
    return ParsedDocument(blocks)


PARSERS = {'pdf': parse_pdf, 'docx': parse_docx, 'txt': parse_txt, 'md': parse_markdown}


def parse_document(data, file_type):
    parsed = PARSERS[file_type](data)
    if not any(block.kind != 'heading' for block in parsed.blocks):
        raise ParseError('文档中没有可入库的正文内容')
    return parsed
