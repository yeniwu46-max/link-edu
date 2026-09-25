import io
import sys
from pathlib import Path

import numpy as np
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app  # noqa: E402
from config import Config  # noqa: E402
from extensions import db  # noqa: E402
from flask_jwt_extended import create_access_token  # noqa: E402
from models import User  # noqa: E402
from rag import embeddings as embedding_module  # noqa: E402
from rag.chunking import Chunker  # noqa: E402
from rag.embeddings import EmbeddingError, HashEmbedder, OpenAICompatibleEmbedder  # noqa: E402
from rag.parsing import parse_docx, parse_markdown, parse_txt, pdf_pages_to_blocks  # noqa: E402
from rag.service import KnowledgeBase, get_kb  # noqa: E402
from rag.settings import INSUFFICIENT_MESSAGE, load_settings  # noqa: E402
from werkzeug.security import generate_password_hash  # noqa: E402

RUBRIC_MD = '''# 微格教学提问技能评价量表

## 一、评价维度

### 提问设计
问题应指向教学目标，难度适中，能够引发学生思考。教师提问后应留出三到五秒的候答时间，避免自问自答。

### 追问与反馈
教师应根据学生回答进行追问，引导学生说明理由；对错误回答给予建设性反馈，而不是直接否定。

## 二、评分标准

| 维度 | 优秀 | 合格 |
| --- | --- | --- |
| 提问设计 | 问题层次清晰 | 问题基本清楚 |
| 候答时间 | 等待三秒以上 | 偶有等待 |
'''

THEORY_TXT = '''第一章 导入技能
导入是课堂教学的起始环节，良好的导入能够激发学习兴趣，建立新旧知识的联系。
常见的导入方式包括直接导入、情境导入、问题导入和实验导入。
第二章 结束技能
结束环节应帮助学生归纳本课要点，形成知识结构，并布置有针对性的练习。
'''


class FakeLLM:
    def __init__(self, response=None):
        self.calls = []
        self.response = response

    def complete_json(self, system, payload, max_tokens):
        self.calls.append(payload)
        return self.response(payload) if callable(self.response) else self.response


@pytest.fixture
def app(tmp_path):
    class TestConfig(Config):
        TESTING = True
        SQLALCHEMY_DATABASE_URI = 'sqlite://'
        JWT_SECRET_KEY = 'test-secret-for-rag-knowledge-base-000000'
        SEED_ON_STARTUP = False
        PUBLIC_DEPLOYMENT = False
        MAX_CONTENT_LENGTH = 4000
        RAG_ENABLED = True
        RAG_STORAGE_DIR = str(tmp_path / 'rag')
        RAG_EMBEDDING_PROVIDER = 'local'
        RAG_ASYNC_INGEST = False

    application = create_app(TestConfig)
    with application.app_context():
        db.create_all()
        for account, role in (('teacher', 'teacher'), ('student', 'student'), ('admin', 'teacher')):
            db.session.add(User(account=account, password_hash=generate_password_hash('x'), name=account, role=role))
        db.session.commit()
        yield application
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    return app.test_client()


def auth(account):
    user = User.query.filter_by(account=account).one()
    return {'Authorization': f'Bearer {create_access_token(identity=str(user.id))}'}


def upload(client, content, filename, account='teacher', **fields):
    data = {'file': (io.BytesIO(content.encode('utf-8') if isinstance(content, str) else content), filename)}
    data.update(fields)
    return client.post('/api/rag/documents', data=data, headers=auth(account), content_type='multipart/form-data')


# ---------- parsing & chunking ----------

def test_markdown_chunks_follow_headings_and_keep_tables():
    chunks = Chunker().split(parse_markdown(RUBRIC_MD.encode()).blocks)
    sections = {chunk.section: chunk for chunk in chunks}
    assert '微格教学提问技能评价量表 > 一、评价维度 > 提问设计' in sections
    assert '候答时间' in sections['微格教学提问技能评价量表 > 一、评价维度 > 提问设计'].text
    table = [c for c in chunks if c.kind == 'table'][0]
    assert table.text.splitlines()[0] == '维度 | 优秀 | 合格'
    assert table.heading_path[-1] == '二、评分标准'


def test_long_paragraph_splits_on_sentence_boundaries_with_overlap():
    paragraph = ''.join(f'第{i}句教师应当关注学生反应并调整节奏。' for i in range(40))
    chunks = Chunker(target_chars=100, max_chars=150, overlap_chars=30, min_chars=0).split(
        parse_txt(paragraph.encode()).blocks)
    assert len(chunks) > 3
    assert all(len(c.text) <= 150 and c.text.endswith('。') for c in chunks)
    assert chunks[1].text.split('\n')[0] == chunks[0].text.split('\n')[-1][-len(chunks[1].overlap):]


def test_short_numbered_items_are_kept_as_content():
    text = '一、评价要点\n1. 目标明确\n2. 重点突出\n3. 方法得当\n'
    chunks = Chunker().split(parse_txt(text.encode()).blocks)
    assert len(chunks) == 1
    assert chunks[0].heading_path == ['一、评价要点']
    assert all(item in chunks[0].text for item in ('目标明确', '重点突出', '方法得当'))


def test_pdf_page_furniture_removed_and_paragraph_spans_pages():
    pages = [
        '微格教学实训规范\n第一章 总则\n微格教学是一种利用现代教学技术手段来培训师范生\n- 1 -',
        '微格教学实训规范\n教学技能的实践方法，每次训练时长为五到十分钟。\n第二章 训练要求\n- 2 -',
        '微格教学实训规范\n训练前应提交教案，训练后进行回放与反思。\n- 3 -',
    ]
    blocks = pdf_pages_to_blocks(pages)
    texts = [b.text for b in blocks]
    assert not any('实训规范' in t or t.strip('- ').isdigit() for t in texts)
    body = next(b for b in blocks if b.text.startswith('微格教学是'))
    assert body.text.endswith('五到十分钟。') and (body.page, body.page_end) == (1, 2)
    assert [b.text for b in blocks if b.kind == 'heading'] == ['第一章 总则', '第二章 训练要求']


def test_cleaning_joins_letter_spacing_only():
    from rag.cleaning import clean_line
    assert clean_line('教 学 评 价\u3000量表\u200b') == '教学评价 量表'
    assert clean_line('第一章  导入技能') == '第一章 导入技能'


def test_docx_styles_and_tables():
    from docx import Document
    document = Document()
    document.add_heading('课堂观察量表', level=1)
    document.add_paragraph('观察者应记录教师提问次数与学生回应情况。')
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = '指标', '描述'
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = '提问', '问题指向目标'
    buffer = io.BytesIO()
    document.save(buffer)
    blocks = parse_docx(buffer.getvalue()).blocks
    assert [(b.kind, b.text.split('\n')[0]) for b in blocks] == [
        ('heading', '课堂观察量表'), ('text', '观察者应记录教师提问次数与学生回应情况。'), ('table', '指标 | 描述')]


# ---------- API flow ----------

def test_upload_retrieve_and_list_chunks(client):
    response = upload(client, RUBRIC_MD, '提问量表.md', category='rubric', tags='提问,量表', source='校内微格中心')
    assert response.status_code == 201, response.json
    document = response.json['document']
    assert document['status'] == 'ready' and document['chunk_count'] >= 3
    assert document['category_label'] == '教学评价量表' and document['tags'] == ['提问', '量表']

    chunks = client.get(f"/api/rag/documents/{document['id']}/chunks", headers=auth('student')).json
    assert chunks['total'] == document['chunk_count'] and chunks['items'][0]['section']

    result = client.post('/api/rag/retrieve', json={'query': '教师提问后应该等待多久的候答时间'},
                         headers=auth('student')).json
    top = result['hits'][0]
    assert top['relevant'] and any('提问设计' in h['section'] for h in result['hits'][:3])
    assert top['document']['title'] == '提问量表' and 0 < top['similarity'] <= 1
    assert result['retrieval']['semantic_embedding'] is False


def test_upload_uses_own_size_limit(client):
    big = RUBRIC_MD + '\n'.join(f'补充说明第{i}条：课堂提问应当面向全体学生。' for i in range(200))
    assert len(big.encode()) > 4000
    assert upload(client, big, 'big.md').status_code == 201


def test_query_answer_keeps_only_valid_citations(app, client):
    upload(client, RUBRIC_MD, '提问量表.md', category='rubric')
    fake = FakeLLM({'sufficient': True, 'answer': '提问后应留出三到五秒候答时间[1][9]。', 'citations': [1, 9]})
    get_kb().llm = fake
    body = client.post('/api/rag/query', json={'query': '提问后候答时间多长'}, headers=auth('student')).json
    assert body['status'] == 'answered' and body['grounded'] is True
    assert body['answer'] == '提问后应留出三到五秒候答时间[1]。'
    assert [c['ref'] for c in body['citations']] == [1]
    assert body['citations'][0]['document_title'] == '提问量表' and body['citations'][0]['section']
    sent = fake.calls[0]['knowledge']
    assert sent[0]['ref'] == 1 and sent[0]['category'] == '教学评价量表'


def test_query_without_relevant_evidence_skips_llm(app, client):
    upload(client, RUBRIC_MD, '提问量表.md', category='rubric')
    fake = FakeLLM({'sufficient': True, 'answer': 'x[1]', 'citations': [1]})
    get_kb().llm = fake
    body = client.post('/api/rag/query', json={'query': 'quantum chromodynamics lattice'}, headers=auth('student')).json
    assert body['status'] == 'insufficient_evidence' and body['answer'].startswith(INSUFFICIENT_MESSAGE)
    assert fake.calls == [] and body['citations'] == []


@pytest.mark.parametrize('response', [
    {'sufficient': False, 'answer': '', 'citations': [], 'missing': '缺少候答时间的量化标准'},
    {'sufficient': True, 'answer': '应等待五秒。', 'citations': [7]},
])
def test_unsupported_answers_become_insufficient(app, client, response):
    upload(client, RUBRIC_MD, '提问量表.md', category='rubric')
    get_kb().llm = FakeLLM(response)
    body = client.post('/api/rag/query', json={'query': '提问后候答时间'}, headers=auth('student')).json
    assert body['status'] == 'insufficient_evidence' and body['answer'].startswith(INSUFFICIENT_MESSAGE)
    assert body['hits'], 'retrieved snippets stay visible to the teacher'


def test_permissions_duplicates_and_bad_files(app, client):
    assert upload(client, RUBRIC_MD, 'a.md', account='student').status_code == 403
    assert upload(client, RUBRIC_MD, 'a.md').status_code == 201
    duplicate = upload(client, RUBRIC_MD, 'b.md')
    assert duplicate.status_code == 409 and duplicate.json['document']['filename'] == 'a.md'
    assert upload(client, 'x', 'old.doc').status_code == 422
    assert upload(client, RUBRIC_MD + '更多', 'c.md', category='unknown').status_code == 422
    app.config['RAG_ADMIN_ACCOUNTS'] = 'admin'
    app.extensions.pop('rag')
    assert upload(client, THEORY_TXT, 'theory.txt').status_code == 403
    assert upload(client, THEORY_TXT, 'theory.txt', account='admin').status_code == 201


def test_filters_deactivate_and_delete(client):
    rubric = upload(client, RUBRIC_MD, '提问量表.md', category='rubric').json['document']
    theory = upload(client, THEORY_TXT, '导入理论.txt', category='theory').json['document']

    def titles(**extra):
        body = client.post('/api/rag/retrieve', json={'query': '导入环节如何激发兴趣', 'top_k': 10, **extra},
                           headers=auth('student')).json
        return {hit['document']['title'] for hit in body['hits']}

    assert titles(categories=['theory']) == {'导入理论'}
    assert titles(document_ids=[rubric['id']]) == {'提问量表'}
    client.patch(f"/api/rag/documents/{theory['id']}", json={'is_active': False}, headers=auth('teacher'))
    assert '导入理论' not in titles()
    assert client.delete(f"/api/rag/documents/{rubric['id']}", headers=auth('teacher')).status_code == 200
    assert titles() == set()
    assert client.post('/api/rag/retrieve', json={'query': 'x', 'categories': ['bad']},
                       headers=auth('student')).status_code == 400


def test_reindex_after_embedding_change(app, client):
    upload(client, THEORY_TXT, '导入理论.txt', category='theory')
    kb = KnowledgeBase(load_settings(), embedder=HashEmbedder(dim=256))
    app.extensions['rag'] = kb
    assert len(kb.stale_documents()) == 1
    assert kb.retrieve('导入方式')['hits'] == []
    for document in kb.stale_documents():
        kb.reindex(document)
    assert kb.stale_documents() == [] and kb.retrieve('导入方式')['hits']


def test_evaluate_marks_theory_basis_per_indicator(app, client):
    upload(client, RUBRIC_MD, '提问量表.md', category='rubric')
    upload(client, THEORY_TXT, '导入理论.txt', category='theory')
    indicators = [
        {'key': 'questioning', 'label': '提问与候答', 'description': '提问后是否留出候答时间并追问'},
        {'key': 'blackboard', 'label': 'zzzz', 'description': 'qqqq'},
    ]

    def respond(payload):
        refs = {item['key']: item['knowledge_refs'] for item in payload['indicators']}
        return {'summary': '整体良好[1]', 'indicators': [
            {'key': 'questioning', 'classroom_evidence': '教师提问后立即自答', 'judgement': '候答不足',
             'theory_refs': refs['questioning'][:1] + [99], 'suggestion': '提问后等待三秒'},
            {'key': 'blackboard', 'classroom_evidence': '证据不足', 'judgement': '无法判断',
             'theory_refs': [1], 'suggestion': ''},
        ]}

    fake = FakeLLM(respond)
    kb = get_kb()
    kb.llm = fake

    def stub_evidence(indicators, **kwargs):
        pool = [{
            'ref': 1,
            'chunk_id': 1,
            'similarity': 0.91,
            'relevant': True,
            'text': '提问后应留出候答时间。',
            'section': '提问设计',
            'page_start': None,
            'page_end': None,
            'document': {'id': 1, 'title': '提问量表', 'category': 'rubric', 'category_label': '教学评价量表', 'source': None},
        }]
        refs = {item['key']: ([1] if item['key'] == 'questioning' else []) for item in indicators}
        return pool, refs

    kb.evidence_for_indicators = stub_evidence
    body = client.post('/api/rag/evaluate', json={
        'transcript': '老师：这道题怎么做？好，我来说答案。',
        'behavior_observations': [{'label': '背对学生板书', 'at_ms': 42000}],
        'indicators': indicators, 'scene': '新授',
    }, headers=auth('student')).json
    assert body['status'] == 'evaluated'
    results = {r['key']: r for r in body['results']}
    assert results['questioning']['theory_status'] == 'grounded'
    assert all(c['ref'] != 99 for c in results['questioning']['citations'])
    assert results['blackboard']['theory_status'] == 'insufficient_evidence'
    assert results['blackboard']['theory_basis'] == INSUFFICIENT_MESSAGE
    classroom = fake.calls[0]['classroom']
    assert classroom['behavior_observations'][0]['label'] == '背对学生板书'

    evidence = client.post('/api/rag/evaluate', json={'generate': False}, headers=auth('student')).json
    assert evidence['status'] == 'evidence_only' and len(evidence['indicators']) == 6
    assert len(fake.calls) == 1


def test_public_deployment_defaults(app, client):
    app.config.update(PUBLIC_DEPLOYMENT=True, PUBLIC_ORIGINS=['https://demo.example'], RAG_ENABLED=None)
    app.extensions.pop('rag', None)
    assert client.get('/api/rag/status', headers=auth('teacher')).status_code == 404

    app.config['RAG_ENABLED'] = True
    app.config['RAG_ASYNC_INGEST'] = False
    app.extensions.pop('rag')
    assert upload(client, THEORY_TXT, 'theory.txt').status_code == 403, 'self-selected teacher role is not enough'

    app.config['RAG_ADMIN_ACCOUNTS'] = 'admin'
    app.extensions.pop('rag')
    big = THEORY_TXT + '\n'.join(f'补充说明第{i}条：导入环节应当联系学生已有经验。' for i in range(200))
    assert len(big.encode()) > app.config['MAX_CONTENT_LENGTH']
    assert upload(client, big, 'theory.txt', account='admin').status_code == 201


def test_status_endpoint(client):
    upload(client, THEORY_TXT, '导入理论.txt', category='theory')
    body = client.get('/api/rag/status', headers=auth('teacher')).json
    assert body['documents_by_category']['theory'] == 1 and body['indexed_chunks'] >= 2
    assert body['can_manage'] is True and body['embedding']['semantic'] is False


# ---------- providers & stores ----------

def test_openai_compatible_embedder_batches_and_checks_dimensions(monkeypatch):
    calls = []

    class Response:
        status_code = 200

        def __init__(self, inputs, dim):
            self.inputs, self.dim = inputs, dim

        def json(self):
            return {'data': [{'index': i, 'embedding': [float(i + 1)] * self.dim}
                             for i in reversed(range(len(self.inputs)))], 'usage': {'total_tokens': 3}}

    def fake_post(url, json, headers, timeout):
        calls.append((url, json))
        return Response(json['input'], 4)

    monkeypatch.setattr(embedding_module.httpx, 'post', fake_post)
    embedder = OpenAICompatibleEmbedder('http://127.0.0.1:9997/v1', 'RAG_TEST_KEY', 'bge-m3', 4, batch_size=2)
    vectors = embedder.embed_documents(['a', 'b', 'c'])
    assert [len(c[1]['input']) for c in calls] == [2, 1] and calls[0][0].endswith('/embeddings')
    assert vectors.shape == (3, 4) and np.allclose(np.linalg.norm(vectors, axis=1), 1)
    with pytest.raises(EmbeddingError):
        OpenAICompatibleEmbedder('http://127.0.0.1:9997/v1', 'K', 'm', 8).embed_query('a')
    with pytest.raises(EmbeddingError):
        OpenAICompatibleEmbedder('http://example.com/v1', 'K', 'm', 4)


def test_faiss_backend_matches_numpy(app, client):
    pytest.importorskip('faiss')
    from rag.vector_store import FaissVectorStore, NumpyVectorStore
    upload(client, RUBRIC_MD, '提问量表.md', category='rubric')
    upload(client, THEORY_TXT, '导入理论.txt', category='theory')
    embedder = get_kb().embedder
    query = embedder.embed_query('导入环节如何激发兴趣')
    exact = NumpyVectorStore(embedder.signature, embedder.dim).search(query, 5)
    approx = FaissVectorStore(embedder.signature, embedder.dim).search(query, 5)
    assert [h.chunk_id for h in exact] == [h.chunk_id for h in approx]
    assert np.allclose([h.score for h in exact], [h.score for h in approx], atol=1e-5)
