from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "临客LINK_双人完善分工_A-B.xlsx"

NAVY = "14213D"
BLUE = "2F6BFF"
BLUE_LIGHT = "EAF1FF"
VIOLET = "7C3AED"
VIOLET_LIGHT = "F2EAFF"
PINK = "EC4899"
PINK_LIGHT = "FCE7F3"
GREEN = "15803D"
GREEN_LIGHT = "DCFCE7"
AMBER = "B45309"
AMBER_LIGHT = "FEF3C7"
RED = "B91C1C"
RED_LIGHT = "FEE2E2"
INK = "1F2937"
MUTED = "667085"
GRID = "D7DFEA"
WHITE = "FFFFFF"

tasks = [
    ["A-01", "P0", "第1轮｜先统一", "全站视觉基线", "统一色彩、字体、间距、圆角、按钮、卡片、图标和弹层规范；减少首页与内页风格割裂。", "A｜主视觉", "视觉规范 + 公共样式", "首页、工作台、课程、评课、成长、资源、个人中心的同类组件观感一致；不出现临时行内样式。", "1.0天", "先完成，供 B 复用状态组件", "纳入"],
    ["A-02", "P0", "第1轮｜先统一", "首页 / 登录主视觉", "优化首屏层级、视频兜底图、登录卡片、角色切换，以及输入错误、加载、成功提示的视觉反馈。", "A｜主视觉", "首页与登录页定稿", "1440×900、1366×768 下主体完整；视频不可播时仍可登录；错误信息不遮挡按钮。", "1.0天", "B 提供真实登录状态", "纳入"],
    ["A-03", "P0", "第1轮｜先统一", "应用壳层与导航", "统一侧栏、顶部搜索、头像菜单、面包屑、当前页高亮及页面切换节奏。", "A｜主视觉", "导航壳层定稿", "六个主导航 + 个人中心/设置/联系我们层级清楚；无文字挤压、遮挡、跳动。", "0.75天", "B 校验所有路由", "纳入"],
    ["A-04", "P0", "第2轮｜模块精修", "关键状态组件", "设计并落地加载、空数据、接口失败、无权限、禁用、保存成功、重试等统一状态。", "A｜主视觉", "状态组件与样式清单", "所有主页面至少有加载/空/失败三态；按钮状态一眼可辨，文案和颜色一致。", "1.0天", "B 接入真实分支", "纳入"],
    ["A-05", "P1", "第2轮｜模块精修", "工作台信息层级", "重排训练摘要、继续训练、反馈摘要、成长轨迹和快捷入口，降低信息密度。", "A｜主视觉", "工作台视觉稿落地", "首屏能看懂“上次练了什么、下一步做什么”；抽屉/弹层在小屏不溢出。", "0.75天", "B 确认真实字段", "纳入"],
    ["A-06", "P1", "第2轮｜模块精修", "课程中心 / 资源库", "统一筛选栏、课程卡片、详情弹层、资源列表；强化可点击区域、来源和文件状态。", "A｜主视觉", "课程与资源页定稿", "筛选、详情、开始训练、打开/无文件状态视觉明确；长标题和空列表不破版。", "1.0天", "B 提供操作结果和文件状态", "纳入"],
    ["A-07", "P1", "第2轮｜模块精修", "AI 评课 / 成长档案", "统一雷达图、折线图、历史列表、报告正文、生成状态和追问区的层级与色彩。", "A｜主视觉", "数据页视觉定稿", "真实/演示报告标识明确；低数据量、长文、生成中、失败重试均可读。", "1.25天", "B 稳定评课和成长数据", "纳入"],
    ["A-08", "P1", "第2轮｜模块精修", "个人中心 / 设置 / 帮助", "整理表单、徽章、联系人、设置开关和客服窗口；补表单校验提示样式。", "A｜主视觉", "个人与帮助页定稿", "表单对齐；成功/失败提示位置固定；客服窗不遮主操作，键盘焦点可见。", "0.75天", "B 补真实保存与留言结果", "纳入"],
    ["A-09", "P0", "第3轮｜联合收口", "响应式与视觉走查", "覆盖 1920、1440、1366×768、1024 宽度；检查滚动、弹层、图表、长文本、空状态。", "A｜主视觉", "视觉问题清单清零", "P0 页面无横向滚动、遮挡、出屏；缩放 125% 可操作；控制台无资源 404。", "1.0天", "与 B 同步回归", "纳入"],
    ["B-01", "P0", "第1轮｜先打通", "登录 / 注册 / 会话", "补输入校验、错误原因、重复提交保护、令牌失效退出；忘记密码必须弹说明或进入可用流程，不能无响应。", "B｜基础功能", "认证闭环", "登录、注册、退出各连续成功 3 次；错误密码、断网、令牌过期均有真实提示；不保留假登录态。", "1.0天", "A 提供状态样式", "纳入"],
    ["B-02", "P0", "第1轮｜先打通", "路由 / 导航 / 全局搜索", "逐项核对侧栏、头像菜单、返回、深链接、刷新；搜索无结果时给结果页反馈。", "B｜基础功能", "导航检查表", "所有可见入口有去向；刷新受保护页面行为正确；搜索命中/无结果都有响应。", "0.75天", "A 定稿导航态", "纳入"],
    ["B-03", "P0", "第1轮｜先打通", "工作台真实响应", "接口失败不再静默伪装成功；继续训练、快捷入口、热力图、日志、反馈摘要均能点击并返回结果。", "B｜基础功能", "工作台交互闭环", "正常/空数据/接口失败三种情形可验；继续训练能带上正确课程；日志刷新后仍存在。", "1.0天", "后端与种子数据可用", "纳入"],
    ["B-04", "P0", "第1轮｜先打通", "课程中心链路", "保证分类、关键词搜索、详情、来源链接、开始训练和课程进度响应；处理无课程与无效 courseId。", "B｜基础功能", "课程链路闭环", "搜索/筛选结果正确；详情可关闭；开始训练参数正确；异常课程不会白屏。", "0.75天", "不改模拟授课内部逻辑", "纳入"],
    ["B-05", "P0", "第2轮｜模块补齐", "AI 评课外围链路", "补报告历史选择、生成中恢复、失败重试、材料不足、追问成功/失败、登录失效处理；明确真实与演示报告边界。", "B｜基础功能", "评课外围闭环", "已有训练可找到对应报告；生成中离开再回来能恢复；失败可重试；追问有空输入/超时/失败提示。", "1.5天", "核心模拟授课只做接口联调", "纳入"],
    ["B-06", "P1", "第2轮｜模块补齐", "成长档案交互", "让 7日/30日/全部切换真实刷新；热力图日期可查看当天记录；回放列表可进入对应评课。", "B｜基础功能", "成长页闭环", "三个时间范围数据不同且正确；日期点击有详情；记录点击进入正确 feedbackId。", "0.75天", "需统一历史评分口径", "纳入"],
    ["B-07", "P0", "第2轮｜模块补齐", "资源库可打开", "为资源补真实 file_url/详情；打开、下载、外链、无文件、失效链接分别响应，不只显示“演示包暂无文件”。", "B｜基础功能", "可用资源清单", "至少 1 份教案、1 份素材、1 份报告可真实打开；缺文件时说明原因和替代入口。", "1.0天", "A 完善文件状态", "纳入"],
    ["B-08", "P1", "第2轮｜模块补齐", "个人中心 / 设置 / 联系", "保存档案失败时不得显示成功；设置持久化可验证；复制邮箱、留言提交要反馈真实结果。", "B｜基础功能", "资料与设置闭环", "保存后刷新仍保留；断网显示失败；清除偏好生效；空留言不可提交；复制失败不报假成功。", "1.0天", "A 提供表单提示样式", "纳入"],
    ["B-09", "P1", "第2轮｜模块补齐", "帮助中心响应", "FAQ 搜索给无结果提示；智能/人工模式切换明确；本地留言与后台能力边界写清。", "B｜基础功能", "帮助交互闭环", "常见问题可命中；无结果可引导；人工留言能落库或明确告知仅本机保存。", "0.5天", "与个人中心留言复用", "纳入"],
    ["B-10", "P0", "第3轮｜联合收口", "启动、错误处理与回归", "统一启动脚本、端口、健康检查、演示账号、环境配置；跑前端测试/构建和后端测试，补关键页面冒烟。", "B｜基础功能", "可交付启动包 + 回归记录", "新电脑按说明可启动；健康检查通过；主导航逐页可开；无白屏/死链；测试与构建通过。", "1.25天", "与 A 联合走查", "纳入"],
    ["AB-01", "P0", "第3轮｜联合收口", "双人联调与演示彩排", "A 盯视觉一致性，B 盯每次点击后的数据/状态；模拟授课只验入口、退出与结果衔接，不在本表重拆核心逻辑。", "A + B", "问题清单 + 演示录像", "登录→工作台→课程→课堂入口→评课→成长→资源→个人中心连续跑 3 次，外围零阻塞。", "0.75天/人", "核心模拟授课需另行已跑通", "纳入"],
    ["B-L1", "P2", "后续迭代", "文件上传与管理", "教案/素材上传、类型和大小校验、删除权限、存储与下载审计。", "B｜基础功能", "上传管理方案", "不阻塞本轮演示；进入试点前再开发。", "待评估", "需要存储与权限方案", "本轮不做"],
    ["B-L2", "P2", "后续迭代", "找回密码 / 第三方登录", "接真实短信/邮箱找回，评估 Google/Apple 等第三方登录。", "B｜基础功能", "认证扩展方案", "本轮必须有明确占位响应，但不要求接第三方平台。", "待评估", "需要外部账号与回调域名", "本轮不做"],
    ["B-L3", "P2", "后续迭代", "生产化与数据治理", "部署、监控、备份、权限审计、音视频/转写 30 天留存与删除策略。", "B｜基础功能", "上线检查表", "试点上线前完成；本轮只保证本机演示稳定。", "待评估", "需要部署环境与制度确认", "本轮不做"],
]

wb = Workbook()
ws = wb.active
ws.title = "A-B分工"
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A9"

headers = ["编号", "优先级", "阶段", "模块", "具体完善项", "负责人", "交付物", "验收标准", "预估", "协作 / 依赖", "本轮范围"]
widths = [10, 9, 17, 22, 55, 16, 25, 55, 12, 30, 12]

ws.merge_cells("A1:K1")
ws["A1"] = "临客 LINK｜核心模拟授课之外的双人完善分工"
ws["A1"].font = Font(name="Microsoft YaHei", size=20, bold=True, color=WHITE)
ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 38

ws.merge_cells("A2:K2")
ws["A2"] = "边界：模拟授课核心逻辑需另行跑通；本表只拆外围页面、视觉统一、基础响应和交付稳定性。"
ws["A2"].font = Font(name="Microsoft YaHei", size=11, bold=True, color=RED)
ws["A2"].fill = PatternFill("solid", fgColor=RED_LIGHT)
ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 30

summary = [
    ("A｜主视觉", "负责统一设计语言、关键页面精修、状态样式、响应式和最终视觉走查。", VIOLET, VIOLET_LIGHT),
    ("B｜基础功能", "负责所有可见按钮/筛选/保存/打开/跳转有真实响应，补错误态、启动交付与回归。", BLUE, BLUE_LIGHT),
    ("优先顺序", "先清 P0：认证、导航、课程、资源、评课外围、状态响应、启动回归；再做 P1 精修。", AMBER, AMBER_LIGHT),
]
for row, (label, desc, color, fill) in enumerate(summary, 4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.cell(row, 1, label)
    ws.cell(row, 1).font = Font(name="Microsoft YaHei", size=11, bold=True, color=color)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=11)
    ws.cell(row, 3, desc)
    ws.cell(row, 3).font = Font(name="Microsoft YaHei", size=10, color=INK)
    ws.cell(row, 3).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row, 3).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 27

header_row = 8
thin = Side(style="thin", color=GRID)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for col, value in enumerate(headers, 1):
    c = ws.cell(header_row, col, value)
    c.font = Font(name="Microsoft YaHei", size=10, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = border
ws.row_dimensions[header_row].height = 32

for r, values in enumerate(tasks, header_row + 1):
    owner = values[5]
    scope = values[10]
    if owner.startswith("A｜"):
        row_fill = VIOLET_LIGHT
    elif owner.startswith("B｜"):
        row_fill = BLUE_LIGHT
    elif owner == "A + B":
        row_fill = GREEN_LIGHT
    else:
        row_fill = AMBER_LIGHT
    if scope == "本轮不做":
        row_fill = "F3F4F6"
    for col, value in enumerate(values, 1):
        c = ws.cell(r, col, value)
        c.font = Font(name="Microsoft YaHei", size=9.5, color=INK)
        c.fill = PatternFill("solid", fgColor=row_fill)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = border
    ws.cell(r, 1).font = Font(name="Microsoft YaHei", size=9.5, bold=True, color=INK)
    ws.cell(r, 2).font = Font(name="Microsoft YaHei", size=9.5, bold=True, color={"P0": RED, "P1": AMBER, "P2": MUTED}[values[1]])
    ws.cell(r, 5).alignment = Alignment(vertical="top", wrap_text=True)
    ws.cell(r, 8).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 64

end_row = header_row + len(tasks)
table = Table(displayName="ABWorkPlan", ref=f"A{header_row}:K{end_row}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
ws.add_table(table)

for i, width in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

ws.auto_filter.ref = f"A{header_row}:K{end_row}"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = f"1:{header_row}"
ws.page_margins.left = 0.2
ws.page_margins.right = 0.2
ws.page_margins.top = 0.35
ws.page_margins.bottom = 0.35
ws.print_area = f"A1:K{end_row}"

ws.conditional_formatting.add(
    f"B{header_row + 1}:B{end_row}",
    FormulaRule(formula=[f'B{header_row + 1}="P0"'], fill=PatternFill("solid", fgColor=RED_LIGHT)),
)

wb.save(OUT)

check = load_workbook(OUT, read_only=False, data_only=False)
assert check.sheetnames == ["A-B分工"]
assert check["A-B分工"].max_row == end_row
assert check["A-B分工"].max_column == len(headers)
check.close()

print(OUT)
