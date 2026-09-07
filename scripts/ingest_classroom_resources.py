"""Import user-supplied teaching files into ignored local RAG storage. No fine-tuning."""
import argparse
import hashlib
import json
import re
import subprocess
import zipfile
from pathlib import Path
from datetime import datetime, timezone
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PRIVATE = ROOT / 'backend/private'

def archive_name(info):
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode('cp437').decode('gbk')
    except (UnicodeError, LookupError):
        return info.filename

def extract(archive):
    folder = PRIVATE / 'materials' / hashlib.sha256(archive.read_bytes()).hexdigest()[:12]
    folder.mkdir(parents=True, exist_ok=True)
    items = []
    with zipfile.ZipFile(archive) as z:
        if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Archive exceeds 100MB import limit')
        for info in z.infolist():
            name = archive_name(info).replace('\\', '/')
            path = (folder / name).resolve()
            if not path.is_relative_to(folder.resolve()):
                raise ValueError('Unsafe archive path')
            if info.is_dir() or path.suffix.lower() not in ('.doc', '.docx', '.pdf', '.md'):
                continue
            if info.file_size > 25 * 1024 * 1024:
                raise ValueError('Document exceeds 25MB import limit')
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(z.read(info))
            items.append((name, path))
    return items

def paragraphs(path):
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        for page, p in enumerate(PdfReader(path).pages, 1):
            text = p.extract_text() or ''
            if text.strip():
                yield f'第{page}页', text
    elif suffix == '.docx':
        for n, p in enumerate(Document(path).paragraphs, 1):
            if p.text.strip():
                yield f'第{n}段', p.text
    elif suffix == '.doc':
        out = path.with_suffix('.paragraphs.json')
        subprocess.run(['powershell', '-NoProfile', '-File', str(ROOT / 'scripts/extract-legacy-doc.ps1'),
                        '-InputPath', str(path), '-OutputPath', str(out)], check=True, timeout=90, capture_output=True)
        for p in json.loads(out.read_text(encoding='utf-8-sig')):
            yield f'第{p["page"]}页·第{p["paragraph"]}段', p['text']
    else:
        for n, p in enumerate(path.read_text(encoding='utf-8-sig').split('\n\n'), 1):
            if p.strip():
                yield f'第{n}段', p

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('archive', type=Path)
    args = parser.parse_args()
    PRIVATE.mkdir(parents=True, exist_ok=True)
    docs, manifest = [], []
    for name, path in extract(args.archive):
        category = '官方大纲（用户附件，来源归属待核验）' if '官方' in name else '高校案例（用户附件）' if '高校' in name else '经验建议（非评价标准）'
        entry = {'file': name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                 'type': category, 'use': '本地检索，最多5段作为云请求上下文；不训练权重，不公开原文',
                 'training_permission': '尚未确认，不可作为后续训练数据'}
        try:
            chunks = []
            for location, text in paragraphs(path):
                text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', text).strip()
                for offset in range(0, len(text), 700):
                    fragment = text[offset:offset+700]
                    if len(fragment) < 8:
                        continue
                    chunks.append({'id': 'user-' + hashlib.sha256((name+location+str(offset)).encode()).hexdigest()[:16],
                        'title': path.stem, 'type': category, 'source': name, 'location': location + (f'·字符{offset+1}起' if offset else ''),
                        'text': fragment, 'training_permission': entry['training_permission']})
            entry.update(status='imported' if chunks else 'empty_needs_ocr', chunks=len(chunks))
            docs.extend(chunks)
        except Exception as exc:
            entry.update(status='failed', error=type(exc).__name__, chunks=0)
        manifest.append(entry)
        print(json.dumps({k:entry[k] for k in ('file','status','chunks')}, ensure_ascii=True))
    (PRIVATE / 'knowledge.json').write_text(json.dumps(docs, ensure_ascii=False, indent=2), encoding='utf-8')
    (PRIVATE / 'materials-manifest.json').write_text(json.dumps({'imported_at':datetime.now(timezone.utc).isoformat(),
        'archive_name':args.archive.name, 'items':manifest}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'Imported {len(docs)} traceable chunks; original files remain private.')

if __name__ == '__main__':
    main()
